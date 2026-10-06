from docx import Document

from gravai.minutes.docx_writer import EMPTY_NOTE, write_minutes
from gravai.minutes.templates import load_template


def render(tmp_path, data, name="ata.docx"):
    path = write_minutes(
        load_template(data["template"]), data["header"], data["participants"], data["minutes"], str(tmp_path / name)
    )
    return Document(path)


def paragraph_texts(doc):
    return [p.text for p in doc.paragraphs if p.text.strip()]


def table_texts(doc):
    return [[c.text for c in row.cells] for t in doc.tables for row in t.rows]


def test_dell_document_follows_the_pdf_order(tmp_path, meeting, speakers, filled_minutes):
    data = {"template": "dell", "header": meeting["header"], "participants": speakers, "minutes": filled_minutes}
    doc = render(tmp_path, data)
    texts = paragraph_texts(doc)

    assert texts[:2] == ["ATA DE REUNIÃO", "Reunião com Dell Technologies"]
    order = ["PARTICIPANTES", "TÓPICOS DISCUTIDOS", "RESUMO DA REUNIÃO", "DECISÕES", "PLANO DE AÇÃO", "PRÓXIMA REUNIÃO"]
    assert [t for t in texts if t in order] == order  # each heading once, in this order

    rows = table_texts(doc)
    assert ["Projeto", meeting["header"]["projeto"], "Ata nº", meeting["header"]["numero"]] in rows
    assert speakers[:3] in rows and [speakers[3], "", ""] in rows  # three columns, last row padded
    assert ["Nº", "Decisão", "Aprovada por"] in rows
    assert ["ID", "Ação / Entregável", "Responsável", "Prazo", "Status"] in rows
    a3 = filled_minutes["plano_acao"][2]
    assert next(r for r in rows if r and r[0] == "A3")[1:] == [a3["acao"], a3["responsavel"], a3["prazo"], a3["status"]]
    nxt = filled_minutes["proxima_reuniao"]
    assert [nxt["data"], nxt["horario"], nxt["pauta"]] in rows

    bands = [r[0] for r in rows if len(r) == 1]
    assert bands == ["Finalizado", "Em andamento", "Ideias e propostas", "Dúvidas em aberto", "Pendências"]

    footer = doc.sections[0].footer.tables[0].rows[0].cells
    assert (
        footer[0].text
        == f"Dell Technologies · Ata de Reunião · {meeting['header']['projeto']} · Ata gerada automaticamente pelo GravAI"
    )
    assert footer[1].text == "Página "


def test_empty_sections_show_a_note_instead_of_disappearing(tmp_path):
    data = {
        "template": "professores",
        "header": {"projeto": "GravAI Atas", "data": "29/09/2026", "professores": "Prof. Marcos"},
        "participants": ["Ana", "Bruno", "Prof. Marcos"],
        "minutes": {
            "topicos": [{"titulo": "Andamento", "resumo": "Ana apresentou o normalizador."}],
            "resumo": "Reunião de orientação.",
            "feedback_professores": ["Prof. Marcos: documentar cada filtro com exemplo real."],
            "decisoes": [],
            "plano_acao": [],
            "proxima_reuniao": {},
        },
    }
    doc = render(tmp_path, data)
    texts = paragraph_texts(doc)
    rows = table_texts(doc)

    assert texts.count(EMPTY_NOTE) == 4  # finalizado, em_andamento, duvidas, pendencias
    assert [r for r in rows if r[0] == EMPTY_NOTE]  # empty tables keep their header row plus the note
    assert ["Professor(es)", "Prof. Marcos", "Tipo", ""] in rows
    assert "ATA DE REUNIÃO COM PROFESSORES" in texts


def test_daily_uses_inline_participants_and_one_group_heading(tmp_path):
    data = {
        "template": "daily",
        "header": {"projeto": "GravAI Atas", "data": "29/09/2026 · 09:00 – 09:12"},
        "participants": ["Ana", "Bruno", "Diego"],
        "minutes": {
            "feito": ["Bruno: normalizador e testes."],
            "em_andamento": ["Bruno: gerador de docx.", "Ana: contrato da API."],
            "ideias": ["Ollama no docker compose (Bruno)."],
            "duvidas": ["Modelo de 8B gera JSON confiável?"],
            "bloqueios": [],
            "combinados": [
                {
                    "id": "A1",
                    "o_que": "Subir o Ollama localmente",
                    "quem": "Diego",
                    "ate_quando": "Próxima daily",
                    "status": "Pendente",
                }
            ],
        },
    }
    doc = render(tmp_path, data)
    texts = paragraph_texts(doc)

    assert "Quem estava: Ana, Bruno, Diego" in texts
    assert texts.count("Como foi") == 1
    assert (
        texts.index("Como foi")
        < texts.index("Feito desde a última daily")
        < texts.index("Travado / bloqueios")
        < texts.index("Combinados")
    )
    assert ["A1", "Subir o Ollama localmente", "Diego", "Próxima daily", "Pendente"] in table_texts(doc)
